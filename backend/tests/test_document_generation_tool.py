"""DocumentGenerationTool tests: validation, workspace isolation, artifact
registration, failure cleanup, sources, resource scheduling, and log hygiene."""

import asyncio
import logging
from pathlib import Path

import pytest

from app.schemas.artifact import ArtifactStatus
from app.schemas.resources import ResourceCapacity, ResourceRequirements
from app.services.artifact_store import InMemoryArtifactStore
from app.services.document_generator import (
    DocumentGenerationError,
    DocumentGenerator,
    WordDocumentGenerator,
    XlsxDocumentGenerator,
)
from app.services.log_context import set_job_context
from app.services.resource_provider import InMemoryResourceProvider
from app.services.resource_scheduler import InMemoryResourceScheduler
from app.services.tool_registry import ToolRegistry
from app.services.tools import DocumentGenerationTool, ToolError
from tests.conftest import default_capacity, make_png

VALID_ARGS = {
    "type": "word",
    "filename": "approval_note.docx",
    "title": "Inspection Approval Note",
    "document_type": "approval_note",
    "sections": [
        {"heading": "Inspection Summary", "paragraphs": ["Pump inspected on 2026-08-15."]},
        {"heading": "Key Findings", "bullets": ["Vibration 2.1 mm/s", "Seal leakage 3 ml/hr"]},
        {"heading": "Required Actions", "numbered": ["Monitor the seal"]},
    ],
    "sources": ["inspection_report.pdf, page 1", "pump_maintenance_manual.txt"],
}


def run(coro):
    return asyncio.run(coro)


def make_tool(tmp_path, requirements=None, capacity=None, generator=None, wait_rounds=3):
    store = InMemoryArtifactStore()
    scheduler = InMemoryResourceScheduler(
        InMemoryResourceProvider(capacity or default_capacity())
    )
    common = {
        "artifact_store": store,
        "scheduler": scheduler,
        "requirements": requirements,
        "wait_rounds": wait_rounds,
    }
    if generator is not None:
        tool = DocumentGenerationTool(generator=generator, **common)
    else:
        # Mirror production wiring: one tool fronting the word + excel generators.
        tool = DocumentGenerationTool(
            generators={
                "word": WordDocumentGenerator(),
                "excel": XlsxDocumentGenerator(),
            },
            **common,
        )
    return tool, store, scheduler


def docx_texts(path):
    from docx import Document

    return [p.text for p in Document(str(path)).paragraphs]


class FailingGenerator(DocumentGenerator):
    supported_types = ("word",)

    async def generate(self, content, output_dir, filename):
        (Path(output_dir) / filename).write_bytes(b"partial data")
        raise DocumentGenerationError("boom")

    def describe(self):
        return {"generator": "failing", "types": ["word"]}


def test_tool_generates_word_artifact(tmp_path):
    tool, store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-t1")
    result = run(tool.execute(tmp_path, VALID_ARGS))
    assert result.ok
    assert "approval_note.docx" in result.content
    assert "Artifact ID:" in result.content

    artifact_file = tmp_path / "artifacts" / "approval_note.docx"
    assert artifact_file.is_file()
    assert artifact_file.stat().st_size > 0
    assert "Inspection Approval Note" in docx_texts(artifact_file)

    artifacts = run(store.list_for_job("job-t1"))
    assert len(artifacts) == 1
    artifact = artifacts[0]
    assert artifact.status == ArtifactStatus.COMPLETED
    assert artifact.type == "word"
    assert artifact.filename == "approval_note.docx"
    assert artifact.size_bytes == artifact_file.stat().st_size
    assert artifact.job_id == "job-t1"
    assert artifact.user_id == "user-001"


def test_tool_artifact_metadata(tmp_path):
    tool, store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-t1")
    run(tool.execute(tmp_path, VALID_ARGS))
    artifacts = run(store.list_for_job("job-t1"))
    assert len(artifacts) == 1
    artifact = artifacts[0]
    assert artifact.status == ArtifactStatus.COMPLETED
    assert artifact.type == "word"
    assert artifact.job_id == "job-t1"
    assert artifact.user_id == "user-001"
    assert artifact.filename == "approval_note.docx"
    assert artifact.size_bytes > 0
    assert artifact.path.endswith("approval_note.docx")


def test_tool_supports_content_string_sections(tmp_path):
    tool, _store, _scheduler = make_tool(tmp_path)
    args = {
        "type": "word",
        "filename": "summary.docx",
        "title": "Summary",
        "sections": [{"heading": "Findings", "content": "Seal leakage 3 ml/hr observed."}],
    }
    set_job_context(user_id="user-001", job_id="job-t2")
    result = run(tool.execute(tmp_path, args))
    assert result.ok
    assert "Seal leakage 3 ml/hr observed." in docx_texts(tmp_path / "artifacts" / "summary.docx")


def test_tool_rejects_unsupported_type(tmp_path):
    # `excel` is now a supported type (see test_tool_generates_excel_artifact);
    # `pdf` remains unsupported.
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-t3")
    with pytest.raises(ToolError, match="unsupported document type"):
        run(tool.execute(tmp_path, {**VALID_ARGS, "type": "pdf"}))
    with pytest.raises(ToolError, match="supported: excel, word"):
        run(tool.execute(tmp_path, {**VALID_ARGS, "type": "pdf"}))


def test_tool_rejects_invalid_filenames(tmp_path):
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-t4")
    for bad in ("../evil.docx", "/abs/path.docx", "dir\\evil.docx", ".hidden.docx", "note.doc", "note"):
        with pytest.raises(ToolError):
            run(tool.execute(tmp_path, {**VALID_ARGS, "filename": bad}))


def test_tool_rejects_malformed_sections(tmp_path):
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-t5")
    with pytest.raises(ToolError, match="unknown section field"):
        run(tool.execute(tmp_path, {**VALID_ARGS, "sections": [{"heading": "X", "bogus": 1}]}))
    with pytest.raises(ToolError, match="paragraphs"):
        run(tool.execute(tmp_path, {**VALID_ARGS, "sections": [{"paragraphs": "not-a-list"}]}))
    with pytest.raises(ToolError, match="at least one"):
        run(tool.execute(tmp_path, {**VALID_ARGS, "sections": [{"heading": ""}]}))


def test_tool_rejects_oversized_content(tmp_path):
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-t6")
    huge = {"heading": "Huge", "paragraphs": ["x" * 200_001]}
    with pytest.raises(ToolError, match="maximum"):
        run(tool.execute(tmp_path, {**VALID_ARGS, "sections": [huge]}))


def test_tool_requires_user_and_job_context(tmp_path):
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context()
    with pytest.raises(ToolError, match="user context"):
        run(tool.execute(tmp_path, VALID_ARGS))
    set_job_context(user_id="user-001")
    with pytest.raises(ToolError, match="job context"):
        run(tool.execute(tmp_path, VALID_ARGS))


def test_tool_writes_only_inside_workspace(tmp_path):
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-t7")
    run(tool.execute(tmp_path, VALID_ARGS))
    artifact_file = tmp_path / "artifacts" / "approval_note.docx"
    assert artifact_file.is_file()
    # nothing outside the workspace or the artifacts subdir
    assert set(tmp_path.iterdir()) == {tmp_path / "artifacts"}
    assert set((tmp_path / "artifacts").iterdir()) == {artifact_file}


def test_generation_failure_cleans_partial_file(tmp_path):
    tool, store, _scheduler = make_tool(tmp_path, generator=FailingGenerator())
    set_job_context(user_id="user-001", job_id="job-t8")
    with pytest.raises(ToolError, match="document_generation failed"):
        run(tool.execute(tmp_path, VALID_ARGS))
    assert not (tmp_path / "artifacts" / "approval_note.docx").exists()
    artifacts = run(store.list_for_job("job-t8"))
    assert len(artifacts) == 1
    assert artifacts[0].status == ArtifactStatus.FAILED


def test_tool_sources_preserved_in_document(tmp_path):
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-t9")
    run(tool.execute(tmp_path, VALID_ARGS))
    texts = docx_texts(tmp_path / "artifacts" / "approval_note.docx")
    assert "inspection_report.pdf, page 1" in texts
    assert "pump_maintenance_manual.txt" in texts


def test_registry_validates_document_generation_args(tmp_path):
    tool, _store, _scheduler = make_tool(tmp_path)
    registry = ToolRegistry([tool])
    set_job_context(user_id="user-001", job_id="job-reg")
    result = run(registry.execute("document_generation", VALID_ARGS, tmp_path))
    assert result.ok
    with pytest.raises(ToolError, match="array of objects"):
        run(registry.execute("document_generation", {**VALID_ARGS, "sections": ["nope"]}, tmp_path))
    with pytest.raises(ToolError, match="array of strings"):
        run(registry.execute("document_generation", {**VALID_ARGS, "sources": [1, 2]}, tmp_path))


def test_resource_scheduling_grant_and_release(tmp_path):
    tool, _store, scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-res")
    result = run(tool.execute(tmp_path, VALID_ARGS))
    assert result.ok
    assert scheduler.stats()["running_jobs"] == 0
    assert scheduler.provider().allocated() == []


def test_resource_scheduling_rejection_fails_cleanly(tmp_path):
    capacity = ResourceCapacity(cpu_cores=0.5, memory_mb=100, gpus=[])
    tool, store, _scheduler = make_tool(tmp_path, capacity=capacity)
    set_job_context(user_id="user-001", job_id="job-reject")
    with pytest.raises(ToolError, match="resources rejected"):
        run(tool.execute(tmp_path, VALID_ARGS))
    artifacts = run(store.list_for_job("job-reject"))
    assert artifacts and artifacts[0].status == ArtifactStatus.FAILED


def test_no_sensitive_content_in_logs(tmp_path, caplog):
    marker = "TOP-SECRET-CONTENT-99"
    args = {
        "type": "word",
        "filename": "secret_note.docx",
        "title": "Secret",
        "sections": [{"heading": "Hidden", "content": marker + " classified findings"}],
    }
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-log")
    with caplog.at_level(logging.INFO, logger="app"):
        result = run(tool.execute(tmp_path, args))
    assert result.ok
    record_text = " ".join(
        str(v) for r in caplog.records for v in r.__dict__.values()
        if not str(v).startswith("<") and "LogRecord" not in str(type(v))
    )
    assert marker not in record_text
    assert "classified findings" not in record_text
    assert "document_generation_started" in record_text
    assert "document_generation_completed" in record_text
    assert "artifact_created" in record_text


# ------------------------------------------------------------------ excel

EXCEL_ARGS = {
    "type": "excel",
    "filename": "findings.xlsx",
    "title": "Inspection Findings",
    "document_type": "spreadsheet",
    "sections": [
        {
            "heading": "Readings",
            "table": [
                ["Item", "Value"],
                ["Course 2", "10.9"],
                ["Course 3", "11.2"],
                ["Total", "=SUM(B2:B3)"],
            ],
        }
    ],
    "sources": ["inspection_report.pdf, page 1"],
}


def test_tool_generates_excel_artifact(tmp_path):
    tool, store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-x1")
    result = run(tool.execute(tmp_path, EXCEL_ARGS))
    assert result.ok

    artifact_file = tmp_path / "artifacts" / "findings.xlsx"
    assert artifact_file.is_file()
    artifacts = run(store.list_for_job("job-x1"))
    assert artifacts[0].type == "excel"
    assert artifacts[0].status == ArtifactStatus.COMPLETED

    from openpyxl import load_workbook

    sheet = load_workbook(artifact_file)["Readings"]
    assert sheet["B2"].value == 10.9
    assert sheet["B4"].value == "=SUM(B2:B3)"
    assert sheet["B4"].data_type == "f"


def test_tool_rejects_excel_extension_mismatch(tmp_path):
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-x2")
    with pytest.raises(ToolError, match=".xlsx"):
        run(tool.execute(tmp_path, {**EXCEL_ARGS, "filename": "findings.docx"}))


# --------------------------------------------------------- approval notes

APPROVAL_ARGS = {
    **VALID_ARGS,
    "approval": {
        "reference_number": "MRPL/OPS/2026/014",
        "date": "2026-08-16",
        "originator": "A. Kumar",
        "department": "Mechanical Maintenance",
        "subject": "Seal replacement approval - cooling water pump",
        "background": "Vibration and seal leakage were observed during inspection.",
        "recommendation": "Approve seal replacement at the next shutdown.",
        "signatures": [
            {"name": "R. Nair", "designation": "Plant Manager", "date": "2026-08-17"},
        ],
    },
}


def test_tool_approval_note_renders_fields(tmp_path):
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-a1")
    result = run(tool.execute(tmp_path, APPROVAL_ARGS))
    assert result.ok

    from docx import Document

    doc = Document(str(tmp_path / "artifacts" / "approval_note.docx"))
    texts = [paragraph.text for paragraph in doc.paragraphs]
    for heading in ("Background", "Recommendation", "Approval"):
        assert heading in texts
    # the title already carries "Approval Note"; there is no duplicate heading
    assert texts[0] == "Inspection Approval Note"
    assert "Approval Note" not in texts
    table_cells = [
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    ]
    assert "MRPL/OPS/2026/014" in table_cells
    assert "Mechanical Maintenance" in table_cells
    # A model-supplied signer name/date is never trustworthy (it cannot know
    # who will actually sign) and must not appear on the printed document.
    assert "R. Nair" not in table_cells
    assert "2026-08-17" not in table_cells
    assert "Plant Manager" in table_cells


def test_tool_never_prints_a_model_supplied_signer_identity(tmp_path):
    """A model has no way to know who will actually approve a note; a name or
    signing date it invents (e.g. "John Doe") must never reach the printed
    document, only the designation/role actually being asked to sign."""
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-a3")
    args = {
        **VALID_ARGS,
        "filename": "approval_note2.docx",
        "approval": {
            **APPROVAL_ARGS["approval"],
            "signatures": [
                {"name": "John Doe", "designation": "Plant Manager", "date": "2026-01-01"},
            ],
        },
    }
    result = run(tool.execute(tmp_path, args))
    assert result.ok

    from docx import Document

    doc = Document(str(tmp_path / "artifacts" / "approval_note2.docx"))
    table_cells = [
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    ]
    assert "John Doe" not in table_cells
    assert "2026-01-01" not in table_cells
    assert "Plant Manager" in table_cells


def test_tool_rejects_malformed_approval(tmp_path):
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-a2")
    with pytest.raises(ToolError, match="unknown approval field"):
        run(tool.execute(tmp_path, {**APPROVAL_ARGS, "approval": {"bogus": 1}}))
    with pytest.raises(ToolError, match="must be a string"):
        run(tool.execute(tmp_path, {**APPROVAL_ARGS, "approval": {"date": 2026}}))
    with pytest.raises(ToolError, match="signatures"):
        run(tool.execute(tmp_path, {**APPROVAL_ARGS, "approval": {"signatures": "none"}}))


# ----------------------------------------------------------------- images


def test_tool_embeds_workspace_image(tmp_path):
    make_png(tmp_path / "crop.png", ["P&ID detail"])
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-i1")
    args = {
        "type": "word",
        "filename": "with_image.docx",
        "title": "Citation Crop",
        "sections": [
            {
                "heading": "Evidence",
                "images": [
                    {"path": "crop.png", "caption": "P&ID citation crop", "width_inches": 4}
                ],
            }
        ],
    }
    result = run(tool.execute(tmp_path, args))
    assert result.ok

    from docx import Document

    doc = Document(str(tmp_path / "artifacts" / "with_image.docx"))
    assert len(doc.inline_shapes) == 1
    assert any(paragraph.text == "P&ID citation crop" for paragraph in doc.paragraphs)


def test_tool_rejects_image_outside_workspace(tmp_path):
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-i2")
    args = {
        "type": "word",
        "filename": "x.docx",
        "title": "X",
        "sections": [{"images": [{"path": "../secret.png"}]}],
    }
    with pytest.raises(ToolError, match="escapes|Absolute"):
        run(tool.execute(tmp_path, args))


def test_tool_rejects_missing_or_bad_image(tmp_path):
    tool, _store, _scheduler = make_tool(tmp_path)
    set_job_context(user_id="user-001", job_id="job-i3")
    missing = {
        "type": "word",
        "filename": "x.docx",
        "title": "X",
        "sections": [{"images": [{"path": "missing.png"}]}],
    }
    with pytest.raises(ToolError, match="not found"):
        run(tool.execute(tmp_path, missing))

    (tmp_path / "bad.gif").write_bytes(b"GIF89a")
    unsupported = {
        "type": "word",
        "filename": "x.docx",
        "title": "X",
        "sections": [{"images": [{"path": "bad.gif"}]}],
    }
    with pytest.raises(ToolError, match="unsupported image type"):
        run(tool.execute(tmp_path, unsupported))
