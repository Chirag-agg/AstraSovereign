"""Excel generator tests: workbook validity, real formulas, numeric cells,
ID preservation, unique sheet names, sources, text fallback, and determinism."""

import asyncio

from app.schemas.document_content import DocumentContent, DocumentSection
from app.services.document_generator import XlsxDocumentGenerator


def run(coro):
    return asyncio.run(coro)


def load(path):
    from openpyxl import load_workbook

    return load_workbook(str(path))


def test_xlsx_generates_valid_workbook(tmp_path):
    content = DocumentContent(
        title="Findings",
        sections=[
            DocumentSection(
                heading="Readings",
                table=[["Item", "Value"], ["Course 2", "10.9"]],
            )
        ],
    )
    generated = run(XlsxDocumentGenerator().generate(content, tmp_path, "findings.xlsx"))
    assert generated.path.exists()
    assert generated.size_bytes > 0
    assert generated.type == "excel"
    workbook = load(generated.path)
    assert "Readings" in workbook.sheetnames
    assert workbook["Readings"]["B2"].value == 10.9


def test_xlsx_writes_real_formulas(tmp_path):
    content = DocumentContent(
        title="Totals",
        sections=[
            DocumentSection(
                heading="Numbers",
                table=[
                    ["Item", "Value"],
                    ["A", "1"],
                    ["B", "2"],
                    ["Total", "=SUM(B2:B3)"],
                ],
            )
        ],
    )
    generated = run(XlsxDocumentGenerator().generate(content, tmp_path, "totals.xlsx"))
    sheet = load(generated.path)["Numbers"]
    assert sheet["B2"].value == 1
    assert sheet["B3"].value == 2
    assert sheet["B4"].value == "=SUM(B2:B3)"
    assert sheet["B4"].data_type == "f"


def test_xlsx_preserves_ids_with_leading_zeros(tmp_path):
    content = DocumentContent(
        title="Ids",
        sections=[
            DocumentSection(heading="Codes", table=[["Code", "Value"], ["007", "42"]])
        ],
    )
    generated = run(XlsxDocumentGenerator().generate(content, tmp_path, "ids.xlsx"))
    sheet = load(generated.path)["Codes"]
    assert sheet["A2"].value == "007"
    assert sheet["B2"].value == 42


def test_xlsx_sheet_names_are_unique(tmp_path):
    content = DocumentContent(
        title="Dup",
        sections=[
            DocumentSection(heading="Data", table=[["a"], ["1"]]),
            DocumentSection(heading="Data", table=[["b"], ["2"]]),
        ],
    )
    generated = run(XlsxDocumentGenerator().generate(content, tmp_path, "dup.xlsx"))
    names = load(generated.path).sheetnames
    assert len(names) == len(set(names)) == 2


def test_xlsx_sources_sheet(tmp_path):
    content = DocumentContent(
        title="Src",
        sections=[DocumentSection(heading="T", table=[["a"], ["1"]])],
        sources=["report.pdf, p.1"],
    )
    generated = run(XlsxDocumentGenerator().generate(content, tmp_path, "src.xlsx"))
    workbook = load(generated.path)
    assert "Sources" in workbook.sheetnames
    sheet = workbook["Sources"]
    assert sheet["A1"].value == "Reference"
    assert sheet["A2"].value == "report.pdf, p.1"


def test_xlsx_formats_measured_values_to_one_decimal(tmp_path):
    content = DocumentContent(
        title="T",
        sections=[
            DocumentSection(
                heading="Readings",
                table=[
                    ["Location", "Reading (mm)", "Limit (mm)", "Margin (mm)", "Status"],
                    ["Tank 204", "10.9", "12.0", "=B2-C2", '=IF(D2>=0,"PASS","FAIL")'],
                ],
            )
        ],
    )
    generated = run(XlsxDocumentGenerator().generate(content, tmp_path, "t.xlsx"))
    sheet = load(generated.path)["Readings"]
    assert sheet["B2"].number_format == "0.0"  # measured value
    assert sheet["C2"].number_format == "0.0"  # limit value
    assert sheet["D2"].number_format == "0.0"  # numeric formula (margin)
    assert sheet["E2"].number_format != "0.0"  # text formula (IF) stays general


def test_xlsx_sources_sheet_has_header_width_and_freeze(tmp_path):
    content = DocumentContent(
        title="T",
        sections=[DocumentSection(heading="T", table=[["a"], ["1"]])],
        sources=["SOP-09, p.12"],
    )
    generated = run(XlsxDocumentGenerator().generate(content, tmp_path, "src.xlsx"))
    sheet = load(generated.path)["Sources"]
    assert sheet["A1"].value == "Reference"
    assert sheet.column_dimensions["A"].width == 30
    assert sheet.freeze_panes == "A2"


def test_xlsx_text_fallback_without_tables(tmp_path):
    content = DocumentContent(
        title="Notes",
        sections=[
            DocumentSection(
                heading="Summary",
                paragraphs=["line one"],
                bullets=["point"],
            )
        ],
    )
    generated = run(XlsxDocumentGenerator().generate(content, tmp_path, "notes.xlsx"))
    sheet = load(generated.path)["Notes"]
    values = [
        cell.value
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    assert "Summary" in values
    assert "line one" in values
    assert "point" in values


def test_xlsx_is_deterministic(tmp_path):
    content = DocumentContent(
        title="D",
        sections=[DocumentSection(heading="T", table=[["a"], ["1"]])],
    )
    generator = XlsxDocumentGenerator()
    first = run(generator.generate(content, tmp_path, "a.xlsx"))
    second = run(generator.generate(content, tmp_path, "b.xlsx"))
    assert first.path.read_bytes() == second.path.read_bytes()


def test_xlsx_freezes_header_and_sizes_columns(tmp_path):
    content = DocumentContent(
        title="T",
        sections=[
            DocumentSection(
                heading="Readings",
                table=[
                    ["Location", "Reading (mm)", "Limit (mm)", "Margin (mm)"],
                    ["Tank 204", "10.9", "12.0", "=B2-C2"],
                ],
            )
        ],
    )
    generated = run(XlsxDocumentGenerator().generate(content, tmp_path, "t.xlsx"))
    sheet = load(generated.path)["Readings"]
    assert sheet.freeze_panes == "A2"
    assert sheet.column_dimensions["A"].width >= len("Location")
    # formula text is ignored when sizing, so computed columns stay tidy
    assert sheet.column_dimensions["D"].width <= 30


def _image_content(paths, width_inches=4.0):
    from app.schemas.document_content import DocumentImage

    return DocumentContent(
        title="Imaged",
        sections=[
            DocumentSection(
                heading="Evidence",
                table=[["Item", "Value"], ["Course 2", "10.9"]],
                images=[
                    DocumentImage(path=str(path), caption=f"Figure {i}", width_inches=width_inches)
                    for i, path in enumerate(paths, start=1)
                ],
            )
        ],
    )


def _media_parts(path):
    import zipfile

    with zipfile.ZipFile(path) as package:
        return [name for name in package.namelist() if name.startswith("xl/media/")]


def test_xlsx_embeds_section_image(tmp_path):
    from tests.conftest import make_png

    picture = make_png(tmp_path / "crop.png", ["P&ID detail"])
    content = _image_content([picture])
    generated = run(XlsxDocumentGenerator().generate(content, tmp_path, "imaged.xlsx"))

    assert _media_parts(generated.path), "workbook carried no image part"
    sheet = load(generated.path)["Evidence"]
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    # Excel has no caption primitive; the caption is written into a cell.
    assert "Figure 1" in values


def test_xlsx_multiple_images_are_anchored_apart(tmp_path):
    from tests.conftest import make_png

    first = make_png(tmp_path / "one.png", ["one"])
    second = make_png(tmp_path / "two.png", ["two"])
    content = _image_content([first, second])
    generated = run(XlsxDocumentGenerator().generate(content, tmp_path, "two.xlsx"))

    assert len(_media_parts(generated.path)) == 2
    sheet = load(generated.path)["Evidence"]
    anchors = [(image.anchor._from.col, image.anchor._from.row) for image in sheet._images]
    assert len(set(anchors)) == 2, "two pictures were stacked on the same anchor"


def test_xlsx_image_is_deterministic(tmp_path):
    from tests.conftest import make_png

    picture = make_png(tmp_path / "crop.png", ["P&ID detail"])
    content = _image_content([picture])
    generator = XlsxDocumentGenerator()
    first = run(generator.generate(content, tmp_path, "a.xlsx"))
    second = run(generator.generate(content, tmp_path, "b.xlsx"))
    assert first.path.read_bytes() == second.path.read_bytes()


def test_xlsx_embeds_converted_bmp_image(tmp_path):
    """A bmp (not a format Excel embeds directly) becomes a real PNG part."""
    from tests.conftest import make_bmp

    picture = make_bmp(tmp_path / "crop.bmp")
    generated = run(
        XlsxDocumentGenerator().generate(_image_content([picture]), tmp_path, "bmp.xlsx")
    )

    parts = _media_parts(generated.path)
    assert parts, "workbook carried no image part"
    assert parts[0].endswith(".png"), f"converted image kept the wrong name: {parts[0]}"


def test_xlsx_converted_image_is_deterministic(tmp_path):
    from tests.conftest import make_bmp

    picture = make_bmp(tmp_path / "crop.bmp")
    generator = XlsxDocumentGenerator()
    first = run(generator.generate(_image_content([picture]), tmp_path, "a.xlsx"))
    second = run(generator.generate(_image_content([picture]), tmp_path, "b.xlsx"))
    assert first.path.read_bytes() == second.path.read_bytes()
