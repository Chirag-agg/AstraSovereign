"""Shared fixtures for backend tests."""

import asyncio
import hashlib
import inspect
import json
import math
import re
import shutil
import subprocess
import time

import httpx
import pytest
from fastapi.testclient import TestClient

from app.config import Settings
from app.main import create_app
from app.schemas.resources import GpuInfo, ResourceCapacity, ResourceRequirements
from app.services.agent import TASK_MARKER
from app.services.auth_store import InMemoryUserStore
from app.services.document_preparer import DocumentPreparer
from app.services.embedding import EmbeddingProvider
from app.services.knowledge_base import KnowledgeBase
from app.services.model_registry import ModelConfig, ModelRegistry
from app.services.multimodal import MultimodalService
from app.services.ocr_provider import FakeOCRProvider, OCRProvider
from app.services.resource_provider import InMemoryResourceProvider
from app.services.resource_scheduler import InMemoryResourceScheduler
from app.services.sandbox_runner import ExecutionResult, SandboxRunnerError
from app.services.vector_store import JsonVectorStore
from app.services.vision_provider import FakeVisionProvider, VisionProvider

DEFAULT_HEADERS = {"X-User-ID": "user-001"}

SANDBOX_IMAGE = "python:3.12-alpine"


def pytest_configure(config):
    """Fail loudly when the suite runs outside the project venv.

    A global interpreter typically lacks the backend runtime deps (e.g.
    python-docx), which surfaces as a scatter of confusing test failures instead
    of one clear error. Only enforce when a local venv exists, CI is not running,
    and the current interpreter is demonstrably missing those deps -- so a
    prepared alternative environment still works.
    """
    import importlib.util
    import os
    import sys
    from pathlib import Path

    if os.environ.get("CI"):
        return
    backend_dir = Path(__file__).resolve().parents[1]
    venv_python = backend_dir / ".venv" / "Scripts" / "python.exe"
    if not venv_python.exists():
        venv_python = backend_dir / ".venv" / "bin" / "python"
    if not venv_python.exists():
        return
    try:
        using_venv = Path(sys.executable).resolve() == venv_python.resolve()
    except OSError:
        using_venv = False
    if using_venv:
        return
    missing = [
        name for name in ("docx", "openpyxl") if importlib.util.find_spec(name) is None
    ]
    if missing:
        raise pytest.UsageError(
            "Backend test dependencies are missing from this interpreter ("
            + ", ".join(missing)
            + "); run the suite with the project venv:\n"
            r"    backend\.venv\Scripts\python.exe -m pytest"
        )


class FakeEmbeddingProvider(EmbeddingProvider):
    """Deterministic, keyword-overlap embedding provider for tests/demos.

    Vectors encode token presence in a hashed space, so cosine similarity ranks
    chunks that share tokens with the query (e.g. "pump maintenance" matches the
    pump document) without needing a real embedding model.
    """

    DIM = 64

    async def embed_many(self, texts: list[str]) -> list[list[float]]:
        return [self._embed_one(text) for text in texts]

    def _embed_one(self, text: str) -> list[float]:
        vector = [0.0] * self.DIM
        for token in re.findall(r"[a-z0-9]+", (text or "").lower()):
            index = int(hashlib.md5(token.encode("utf-8")).hexdigest(), 16) % self.DIM
            vector[index] += 1.0
        norm = math.sqrt(sum(x * x for x in vector)) or 1.0
        return [x / norm for x in vector]

    def describe(self) -> dict:
        return {"provider": "fake", "model": "keyword-hash"}


class StubClassifier:
    """Deterministic capability classifier for tests.

    A stub, not a fake embedder: it maps a message to a capability by explicit
    override or keyword (mirroring the retired keyword router), so tests
    exercise the node sequence without depending on embedding geometry.
    """

    _KEYWORDS = {
        "coding": (
            "```", "def ", "function ", "class ", "import ", "print(", "console.log",
            "=>", "javascript", "typescript", "python", "java ", "sql ", "select ",
            "insert ", "update ", "delete ", "debug", "bug", "fix this", "write code",
            "write a program", "write some code", "code block", "code snippet",
            "code review", "code", "script", "refactor", "regex", "bash", "html",
            "css", "algorithm",
        ),
        "vision": (
            "image", "photo", "picture", "screenshot", "vision", "see this image",
            "look at this image", "attached image",
        ),
        "document": (
            "docx", "upload this file", "report file", "read this file",
            "read the file", "process this document", "summarize the file",
            "summarize this file", "summarize the document", "summarize this document",
        ),
    }

    def __init__(self, overrides: dict | None = None) -> None:
        self._overrides = dict(overrides or {})
        self.seen: list[str] = []

    async def classify(self, message: str):
        from app.services.capability_classifier import CapabilityClassification

        self.seen.append(message)
        text = (message or "").lower()
        for needle, label in self._overrides.items():
            if needle in text:
                return CapabilityClassification(
                    task_type=label, reason="stub override", confidence=1.0
                )
        for label, keywords in self._KEYWORDS.items():
            if any(keyword in text for keyword in keywords):
                return CapabilityClassification(
                    task_type=label, reason="stub keyword", confidence=1.0
                )
        return CapabilityClassification(
            task_type="general", reason="stub default", confidence=1.0
        )


def default_capacity() -> ResourceCapacity:
    """Default test capacity: 8 cores, 16 GB RAM, one 16 GB GPU."""
    return ResourceCapacity(
        cpu_cores=8.0,
        memory_mb=16384,
        gpus=[GpuInfo(gpu_id="GPU-0", vram_mb=16384)],
    )


def docker_ready() -> bool:
    """True when the Docker daemon is reachable AND the sandbox image exists."""
    if shutil.which("docker") is None:
        return False
    try:
        info = subprocess.run(
            ["docker", "info"], capture_output=True, timeout=10, text=True
        )
        if info.returncode != 0:
            return False
        inspect = subprocess.run(
            ["docker", "image", "inspect", SANDBOX_IMAGE],
            capture_output=True,
            timeout=10,
            text=True,
        )
        return inspect.returncode == 0
    except Exception:
        return False


def node_ready() -> bool:
    """True when the Node.js runtime is available for the real PPTX renderer.

    Deliberately does not check the vendored ``presentation/node_modules`` tree:
    that tree is committed, so its absence is a packaging regression the
    ``node``-marked test must fail on rather than silently skip.
    """
    return shutil.which("node") is not None


class FakeSandboxRunner:
    """Scriptable sandbox runner for deterministic tests (no Docker required)."""

    def __init__(self, results=None, error=None):
        self.results = list(results or [])
        self.error = error
        self.calls = []

    async def run(self, code, language="python", stdin=""):
        self.calls.append({"code": code, "language": language, "stdin": stdin})
        if self.error:
            raise self.error
        if self.results:
            return self.results.pop(0)
        return ExecutionResult(
            success=True, exit_code=0, stdout="", stderr="", timed_out=False, duration_ms=1
        )


def ok_result(stdout="", stderr="", exit_code=0, timed_out=False, duration_ms=10):
    return ExecutionResult(
        success=exit_code == 0 and not timed_out,
        exit_code=exit_code,
        stdout=stdout,
        stderr=stderr,
        timed_out=timed_out,
        duration_ms=duration_ms,
    )


def build_registry(models: dict[str, dict]) -> ModelRegistry:
    """Build a ModelRegistry from a plain-dict mapping for tests."""
    return ModelRegistry(
        models={name: ModelConfig(**cfg) for name, cfg in models.items()}
    )


def make_multimodal_stack(
    tmp_path,
    ocr: OCRProvider = None,
    vision: VisionProvider = None,
    vision_model="vision-model",
    vision_enabled=True,
    vision_resources=None,
    capacity=None,
    vision_wait_rounds=3,
    extraction_store=None,
):
    """Build a fully wired MultimodalService for deterministic tests.

    Returns ``(service, scheduler, uploads_root, tmp_root)``. Pass
    ``extraction_store`` when the test needs to inspect the extraction artifact
    the ingestion seam persisted.
    """
    from pathlib import Path

    uploads = Path(tmp_path) / "uploads"
    tmproot = Path(tmp_path) / "tmp"
    uploads.mkdir(parents=True, exist_ok=True)
    tmproot.mkdir(parents=True, exist_ok=True)
    kb = KnowledgeBase(
        vector_store=JsonVectorStore(str(Path(tmp_path) / "kb")),
        embedding_provider=FakeEmbeddingProvider(),
        extraction_store=extraction_store,
    )
    scheduler = InMemoryResourceScheduler(
        InMemoryResourceProvider(capacity or default_capacity())
    )
    service = MultimodalService(
        knowledge_base=kb,
        preparer=DocumentPreparer(),
        ocr_provider=ocr if ocr is not None else FakeOCRProvider(),
        vision_provider=vision if vision is not None else FakeVisionProvider(),
        vision_model=vision_model,
        vision_enabled=vision_enabled,
        vision_requirements=vision_resources or ResourceRequirements(),
        scheduler=scheduler,
        uploads_root=str(uploads),
        tmp_root=str(tmproot),
        vision_wait_rounds=vision_wait_rounds,
    )
    return service, scheduler, uploads, tmproot


def _extract_task(prompt: str) -> str:
    """Extract the user task from an agent prompt for mock echo replies."""
    if TASK_MARKER in prompt:
        tail = prompt.split(TASK_MARKER, 1)[1].strip()
        if tail:
            return tail.splitlines()[0].strip()
    return prompt


def _extract_chat_task(messages) -> str:
    for message in reversed(messages or []):
        content = message.get("content")
        if content:
            return _extract_task(str(content))
    return ""


def _chat_message_from_script(text: str) -> dict:
    """Translate a scripted JSON envelope into an Ollama native chat message.

    This is the adapter that lets the existing scripted tests exercise the native
    ``/api/chat`` path without rewriting their scripts. Unknown JSON is passed
    through as assistant content so the agent's legacy fallback (and its
    retry-on-garbage behaviour) still applies.
    """
    try:
        parsed = json.loads(text)
    except (ValueError, json.JSONDecodeError):
        parsed = None
    if isinstance(parsed, dict):
        if parsed.get("type") == "final" or (
            isinstance(parsed.get("response"), str) and "tool" not in parsed
        ):
            return {"role": "assistant", "content": parsed.get("response") or ""}
        tool = parsed.get("tool")
        if isinstance(tool, str) and tool:
            arguments = parsed.get("arguments")
            if not isinstance(arguments, dict):
                arguments = {}
            return {
                "role": "assistant",
                "content": "",
                "tool_calls": [{"function": {"name": tool, "arguments": arguments}}],
            }
    return {"role": "assistant", "content": text}


def make_native_chat_handler(turns, arguments_as_string: bool = False):
    """Mock Ollama returning RAW ``/api/chat`` responses in Ollama's real shape.

    ``turns`` is a list of ``{"content": str, "tool_calls": [{"name", "arguments"}]}``.
    ``arguments_as_string`` emits ``arguments`` as a JSON string (a real variant
    some templates produce) instead of an object.
    """

    script = list(turns)

    async def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/api/tags":
            return httpx.Response(200, json={"models": [{"name": "test-model"}]})
        if request.url.path == "/api/chat":
            if not script:
                raise AssertionError("model called more times than scripted")
            turn = script.pop(0)
            message: dict = {"role": "assistant", "content": turn.get("content", "")}
            if turn.get("tool_calls"):
                message["tool_calls"] = [
                    {
                        "function": {
                            "name": call["name"],
                            "arguments": json.dumps(call["arguments"])
                            if arguments_as_string
                            else call["arguments"],
                        }
                    }
                    for call in turn["tool_calls"]
                ]
            return httpx.Response(
                200, json={"model": "test-model", "message": message}
            )
        return httpx.Response(404, json={"error": "not found"})

    return handler


def make_ollama_handler(available_models, delay_seconds: float = 0.0):
    """Mock Ollama: /api/tags lists ``available_models``, /api/generate replies
    only for known models (404 otherwise, matching real Ollama)."""

    available = set(available_models or set())

    async def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/api/tags":
            return httpx.Response(
                200, json={"models": [{"name": name} for name in sorted(available)]}
            )
        if request.url.path == "/api/generate":
            if delay_seconds:
                await asyncio.sleep(delay_seconds)
            payload = json.loads(request.content)
            model = payload.get("model")
            if model not in available:
                return httpx.Response(
                    404, json={"error": f"model '{model}' not found"}
                )
            task = _extract_task(payload["prompt"])
            return httpx.Response(
                200,
                json={
                    "response": f"mocked reply to: {task}",
                    "model": model,
                },
            )
        if request.url.path == "/api/chat":
            if delay_seconds:
                await asyncio.sleep(delay_seconds)
            payload = json.loads(request.content)
            model = payload.get("model")
            if model not in available:
                return httpx.Response(
                    404, json={"error": f"model '{model}' not found"}
                )
            task = _extract_chat_task(payload.get("messages"))
            return httpx.Response(
                200,
                json={
                    "model": model,
                    "message": {
                        "role": "assistant",
                        "content": f"mocked reply to: {task}",
                    },
                },
            )
        return httpx.Response(404, json={"error": "not found"})

    return handler


def make_scripted_handler(responses, delay_seconds: float = 0.0):
    """Mock Ollama that plays back a scripted list of model outputs.

    Calling the model more times than scripted raises an AssertionError, which
    fails the job (used to prove the agent does not exceed its budget).
    """

    script = list(responses)

    async def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/api/tags":
            return httpx.Response(
                200, json={"models": [{"name": "test-model"}, {"name": "coder-model"}]}
            )
        if request.url.path == "/api/generate":
            if delay_seconds:
                await asyncio.sleep(delay_seconds)
            if not script:
                raise AssertionError("model called more times than scripted")
            text = script.pop(0)
            return httpx.Response(200, json={"response": text, "model": "test-model"})
        if request.url.path == "/api/chat":
            if delay_seconds:
                await asyncio.sleep(delay_seconds)
            if not script:
                raise AssertionError("model called more times than scripted")
            text = script.pop(0)
            return httpx.Response(
                200,
                json={"model": "test-model", "message": _chat_message_from_script(text)},
            )
        return httpx.Response(404, json={"error": "not found"})

    return handler


def _wrap_lifecycle_calls(handler):
    """Intercept OllamaService.unload_and_wait's own HTTP traffic (a
    ``/api/generate`` call with no ``prompt`` and ``keep_alive: 0``, plus its
    ``/api/ps`` polling) before it ever reaches a test-authored handler.

    These are model-lifecycle calls, not model turns: a scripted handler's
    response list is a script of *model turns*, and none of the make_*
    handler factories above know about this traffic, so without this they'd
    either consume a scripted response meant for a real turn or raise on an
    unhandled path. Applied once, at the single point (client_factory) where
    every test's handler meets OllamaService's transport.
    """

    async def wrapped(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/api/generate":
            try:
                payload = json.loads(request.content)
            except ValueError:
                payload = {}
            # Unambiguously an unload call, never a legitimate scripted model
            # turn (which always carries a prompt) — always short-circuited.
            if payload.get("keep_alive") == 0 and "prompt" not in payload:
                return httpx.Response(200, json={})
        result = handler(request)
        if inspect.isawaitable(result):
            result = await result
        # A test that defines its own /api/ps response (e.g. to assert on
        # residency) takes priority; only default to "nothing resident" when
        # the handler doesn't know about the path at all (its 404 fallback).
        if request.url.path == "/api/ps" and result.status_code == 404:
            return httpx.Response(200, json={"models": []})
        return result

    return wrapped


def make_mutable_scripted_handler(script, delay_seconds: float = 0.0):
    """Mock Ollama that reads a caller-mutated list of model outputs live.

    Used when the scripted agent replies must be assembled after a document is
    uploaded (e.g. ``document_vision`` needs the generated ``document_id``).
    """

    async def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/api/tags":
            return httpx.Response(
                200,
                json={
                    "models": [
                        {"name": "test-model"},
                        {"name": "coder-model"},
                        {"name": "vision-model"},
                    ]
                },
            )
        if request.url.path == "/api/generate":
            if delay_seconds:
                await asyncio.sleep(delay_seconds)
            if not script:
                raise AssertionError("model called more times than scripted")
            text = script.pop(0)
            return httpx.Response(200, json={"response": text, "model": "test-model"})
        if request.url.path == "/api/chat":
            if delay_seconds:
                await asyncio.sleep(delay_seconds)
            if not script:
                raise AssertionError("model called more times than scripted")
            text = script.pop(0)
            return httpx.Response(
                200,
                json={"model": "test-model", "message": _chat_message_from_script(text)},
            )
        return httpx.Response(404, json={"error": "not found"})

    return handler


def make_text_pdf(path, lines):
    """Generate a text-based PDF with the given lines (reportlab, local only)."""
    from reportlab.pdfgen import canvas

    canvas_obj = canvas.Canvas(str(path))
    y = 760
    for line in lines:
        canvas_obj.drawString(72, y, line)
        y -= 18
    canvas_obj.save()


def make_blank_pdf(path, pages=1):
    """Generate an image-only (no text) PDF — used for OCR-required detection.

    Contains ``pages`` blank pages (no text layer) so it remains renderable by
    pypdfium2 while reporting no extractable text.
    """
    from reportlab.pdfgen import canvas

    canvas_obj = canvas.Canvas(str(path))
    for _ in range(pages):
        canvas_obj.showPage()
    canvas_obj.save()


def make_png(path, lines, width=600, height=220):
    """Generate a PNG image containing the given text lines (Pillow, local only)."""
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    y = 40
    for line in lines:
        draw.text((20, y), line, fill="black")
        y += 50
    image.save(path)
    return path


def make_image_pdf(path, image_path, page_count=1):
    """Embed a raster image into an otherwise text-less PDF (scanned-style)."""
    from PIL import Image
    from reportlab.pdfgen import canvas

    image = Image.open(image_path)
    c = canvas.Canvas(str(path), pagesize=(image.width, image.height))
    for _ in range(page_count):
        c.drawImage(str(image_path), 0, 0, width=image.width, height=image.height)
        c.showPage()
    c.save()


def make_bmp(path, size=(120, 80), color=(200, 30, 30)):
    """A non-web-safe raster, so the embedding path must convert it to PNG."""
    from PIL import Image

    Image.new("RGB", size, color).save(path, format="BMP")
    return path


def make_image_bytes(size=(140, 100), color=(20, 90, 200), fmt="PNG"):
    """In-memory raster bytes, for embedding a picture in a container."""
    import io

    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", size, color).save(buffer, format=fmt)
    return buffer.getvalue()


def make_animated_gif(path, frames=3):
    """An animated GIF: each frame is a page the reader would page through."""
    from PIL import Image

    images = [
        Image.new("RGB", (60, 40), (index * 70 % 256, 40, 40))
        for index in range(frames)
    ]
    images[0].save(
        path, save_all=True, append_images=images[1:], duration=100, loop=0
    )
    return path


def make_multipage_tiff(path, frames=3):
    """A multi-page TIFF — the other multi-frame raster family."""
    from PIL import Image

    images = [
        Image.new("RGB", (60, 40), (40, index * 70 % 256, 40))
        for index in range(frames)
    ]
    images[0].save(path, save_all=True, append_images=images[1:])
    return path


def make_docx(path, paragraphs=(), tables=(), pictures=()):
    """A .docx with the given paragraphs and tables, in document order.

    ``tables`` is a sequence of row lists (each row a list of cell strings).
    Paragraphs come first, then tables — sufficient for the extractor tests.
    ``pictures`` is a sequence of raster byte strings; each is embedded as a
    real inline image (python-docx writes it under ``word/media/``), which is
    how the embedded-picture path sees a container that carries one.
    """
    import io

    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    for rows in tables:
        table = document.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    for data in pictures:
        document.add_picture(io.BytesIO(data))
    document.save(str(path))
    return path


def make_xlsx(path, sheets):
    """An .xlsx from ``{sheet_name: [[cell, ...], ...]}``.

    A cell whose text begins with "=" is stored as a formula (openpyxl's own
    rule), so a sheet written this way has no cached value for that cell.
    """
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    workbook.save(str(path))
    return path


def make_pptx(path, slides, notes=None, media=None):
    """A minimal .pptx built with stdlib zipfile (python-pptx is not installed).

    ``slides`` is a list of ``(title, body_lines)`` pairs; ``notes`` maps a
    1-based slide number to its speaker-notes text. The parts are the minimum
    the extractor reads: ``ppt/slides/slideN.xml`` and, when present,
    ``ppt/notesSlides/notesSlideN.xml``.

    ``media`` maps a 1-based slide number to raster byte strings placed on that
    slide: each is written under ``ppt/media/`` and linked from the slide's
    rels part, which is how a real deck records a picture.
    """
    import zipfile

    def _drawingml(lines):
        paragraphs = "".join(
            f"<a:p><a:r><a:t>{line}</a:t></a:r></a:p>" for line in lines
        )
        return (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
            ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f"<p:cSld><p:spTree><p:sp><p:txBody>{paragraphs}</p:txBody></p:sp>"
            "</p:spTree></p:cSld></p:sld>"
        )

    def _rels(targets):
        relationships = "".join(
            '<Relationship Id="rId{index}" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/image" Target="../media/{name}"/>'.format(
                index=index, name=name
            )
            for index, name in enumerate(targets, start=1)
        )
        return (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/'
            f'2006/relationships">{relationships}</Relationships>'
        )

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for index, (title, body) in enumerate(slides, start=1):
            archive.writestr(
                f"ppt/slides/slide{index}.xml", _drawingml([title, *body])
            )
        for number, text in (notes or {}).items():
            archive.writestr(
                f"ppt/notesSlides/notesSlide{number}.xml", _drawingml([text])
            )
        written: list[str] = []
        for number, pictures in sorted((media or {}).items()):
            targets = []
            for data in pictures:
                written.append(data)
                name = f"image{len(written)}.png"
                archive.writestr(f"ppt/media/{name}", data)
                targets.append(name)
            archive.writestr(f"ppt/slides/_rels/slide{number}.xml.rels", _rels(targets))
    return path


def make_odt(path, paragraphs=(), tables=(), pictures=()):
    """A minimal OpenDocument text file (stdlib zipfile; odfpy is not installed).

    ``pictures`` is a sequence of raster byte strings, each written under
    ``Pictures/`` — where an ODF file keeps the images it embeds.
    """
    import zipfile

    blocks = "".join(
        f'<text:p text:style-name="P1">{text}</text:p>' for text in paragraphs
    )
    for rows in tables:
        cells = "".join(
            "<table:table-row>"
            + "".join(
                f"<table:table-cell><text:p>{value}</text:p></table:table-cell>"
                for value in row
            )
            + "</table:table-row>"
            for row in rows
        )
        blocks += f"<table:table>{cells}</table:table>"
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content'
        ' xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"'
        ' xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">'
        f"<office:body><office:text>{blocks}</office:text></office:body>"
        "</office:document-content>"
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", content)
        for index, data in enumerate(pictures, start=1):
            archive.writestr(f"Pictures/image{index}.png", data)
    return path


def wait_for_job(
    client: TestClient,
    job_id: str,
    user_id: str = "user-001",
    terminal=("completed", "failed", "cancelled"),
    timeout: float = 5.0,
) -> dict:
    """Poll a job until it reaches a terminal state. Used because the worker
    processes jobs asynchronously in the background."""
    deadline = time.monotonic() + timeout
    last_body = None
    while time.monotonic() < deadline:
        resp = client.get(f"/api/jobs/{job_id}", headers={"X-User-ID": user_id})
        assert resp.status_code == 200, resp.text
        last_body = resp.json()
        if last_body["status"] in terminal:
            return last_body
        time.sleep(0.02)
    raise AssertionError(
        f"job {job_id} did not reach a terminal state within {timeout}s "
        f"(last status: {last_body and last_body['status']})"
    )


@pytest.fixture
def app_settings(tmp_path):
    return Settings(
        ollama_base_url="http://ollama.test",
        default_model="test-model",
        log_level="ERROR",
        log_file="",
        workspaces_root=str(tmp_path / "workspaces"),
        knowledge_base_root=str(tmp_path / "knowledge"),
        extraction_root=str(tmp_path / "extractions"),
        uploads_root=str(tmp_path / "uploads"),
        multimodal_tmp_root=str(tmp_path / "tmp"),
        audit_root=str(tmp_path / "audit"),
        job_store_root=str(tmp_path / "jobs"),
        artifact_store_root=str(tmp_path / "artifacts"),
        database_path=str(tmp_path / "astra.db"),
    )


@pytest.fixture
def test_models():
    """Default test registry: general + coding enabled, document/vision disabled."""
    return {
        "general": {
            "provider": "ollama",
            "model": "test-model",
            "enabled": True,
            "capabilities": ["general", "reasoning", "summarization"],
        },
        "coding": {
            "provider": "ollama",
            "model": "coder-model",
            "enabled": True,
            "capabilities": ["coding", "debugging", "code_review"],
        },
        "document": {
            "provider": "ollama",
            "model": "doc-model",
            "enabled": False,
            "capabilities": ["document", "summarization"],
        },
        "vision": {
            "provider": "ollama",
            "model": "vision-model",
            "enabled": False,
            "capabilities": ["vision", "image"],
        },
    }


@pytest.fixture
def available_models():
    return {"test-model", "coder-model"}


@pytest.fixture
def success_ollama_handler(available_models):
    return make_ollama_handler(available_models)


@pytest.fixture
def delayed_ollama_handler(available_models):
    def _make(delay_seconds: float = 0.2):
        return make_ollama_handler(available_models, delay_seconds=delay_seconds)

    return _make


@pytest.fixture
def client_factory(app_settings, test_models):
    def _make(
        handler,
        models=None,
        sandbox_enabled=False,
        sandbox_runner=None,
        resource_capacity=None,
        embedding_provider=None,
        ocr_provider=None,
        vision_provider=None,
        classifier=None,
        project_root=None,
        file_max_bytes=1_000_000,
        presentation_renderer=None,
        seed_users=None,
    ):
        registry = build_registry(models if models is not None else test_models)
        updates = {}
        if sandbox_enabled:
            updates["sandbox_enabled"] = True
        if project_root is not None:
            updates["cowork_projects_root"] = str(project_root)
        updates["cowork_file_max_bytes"] = file_max_bytes
        settings = app_settings.model_copy(update=updates) if updates else app_settings

        # seed_users: [{"username": ..., "password": ..., "role": "user"}, ...].
        # When given, a fresh InMemoryUserStore is seeded with exactly these
        # accounts (skipping the real SqliteUserStore + its unknown-password
        # bootstrap admin), so a test can log in as a known account. Seeding
        # happens now, before the lifespan's own bootstrap-if-empty check
        # runs, so that check correctly finds the store non-empty and skips.
        user_store = None
        if seed_users:
            user_store = InMemoryUserStore()
            for entry in seed_users:
                asyncio.run(
                    user_store.create(
                        entry["username"], entry["password"], role=entry.get("role", "user")
                    )
                )

        app = create_app(
            settings=settings,
            ollama_transport=httpx.MockTransport(_wrap_lifecycle_calls(handler)),
            model_registry=registry,
            sandbox_runner=sandbox_runner,
            resource_capacity=resource_capacity,
            embedding_provider=embedding_provider or FakeEmbeddingProvider(),
            ocr_provider=ocr_provider or FakeOCRProvider(),
            vision_provider=vision_provider,
            classifier=classifier or StubClassifier(),
            presentation_renderer=presentation_renderer,
            user_store=user_store,
            # Every existing test authenticates via X-User-ID/X-Role headers;
            # only this fixture may enable the fallback (see deps.py) —
            # production's bare create_app() never does.
            dev_header_auth=True,
        )
        return TestClient(app)

    return _make


def login(client: TestClient, username: str, password: str) -> dict:
    """Log in via the real auth API; TestClient's cookie jar then carries the
    session automatically on subsequent requests in the same `with` block."""
    response = client.post(
        "/api/auth/login", json={"username": username, "password": password}
    )
    response.raise_for_status()
    return response.json()


@pytest.fixture
def client(client_factory, success_ollama_handler):
    with client_factory(success_ollama_handler) as c:
        yield c
