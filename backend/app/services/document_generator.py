"""Local document generator abstraction (Phase 9).

The Agent produces structured ``DocumentContent``; a concrete generator converts
it into an office file. Word (``.docx``) and Excel (``.xlsx``) are implemented;
PowerPoint is handled by the separate presentation renderer. Everything runs
locally with local Python libraries — no cloud document APIs, no external
conversion services, no telemetry.
"""

import asyncio
import logging
import re
from abc import ABC, abstractmethod
from datetime import datetime
from pathlib import Path
from typing import Optional

from app.schemas.document_content import (
    ApprovalSignature,
    DocumentContent,
    DocumentSection,
    GeneratedDocument,
)
from app.services.image_normalization import ImageNormalizationError, embed_source
from app.services.ooxml import normalize_ooxml

logger = logging.getLogger("app.document_generator")


class DocumentGenerationError(Exception):
    """A document could not be generated or validated."""


class DocumentGenerator(ABC):
    """Converts structured document content into a deliverable file."""

    supported_types: tuple[str, ...] = ()

    @abstractmethod
    async def generate(
        self, content: DocumentContent, output_dir: Path, filename: str
    ) -> GeneratedDocument:
        raise NotImplementedError

    def describe(self) -> dict:
        return {"generator": self.__class__.__name__, "types": list(self.supported_types)}


class WordDocumentGenerator(DocumentGenerator):
    """Generates valid .docx files from structured content using python-docx."""

    supported_types = ("word",)

    async def generate(
        self, content: DocumentContent, output_dir: Path, filename: str
    ) -> GeneratedDocument:
        return await asyncio.to_thread(self._generate_sync, content, output_dir, filename)

    def _generate_sync(
        self, content: DocumentContent, output_dir: Path, filename: str
    ) -> GeneratedDocument:
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentGenerationError(
                "python-docx is not installed (Word generation unavailable)"
            ) from exc

        output_dir.mkdir(parents=True, exist_ok=True)
        target = output_dir / filename

        try:
            doc = Document()
            from docx.shared import Mm

            for page_section in doc.sections:
                page_section.page_width = Mm(210)  # A4
                page_section.page_height = Mm(297)
            if content.title:
                doc.add_heading(content.title, level=0)
            if content.subtitle:
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(content.subtitle)
                run.italic = True

            if content.approval is not None:
                self._add_approval_header(doc, content.approval)

            for section in content.sections:
                self._add_section(doc, section)

            # The recommendation and signature block close an approval note, after
            # any findings sections, so the document reads in the right order.
            if content.approval is not None:
                self._add_approval_footer(doc, content.approval)

            if content.sources:
                doc.add_heading("Sources", level=1)
                for source in content.sources:
                    doc.add_paragraph(source, style="List Number")

            self._set_footer(doc, content.classification)

            doc.save(str(target))
        except DocumentGenerationError:
            raise
        except ImageNormalizationError as exc:
            # The message names the file only (never its resolved location), and
            # is the actionable part for a caller: which image could not be read.
            raise DocumentGenerationError(f"Word generation failed: {exc}") from exc
        except Exception as exc:
            raise DocumentGenerationError(
                f"Word generation failed: {exc.__class__.__name__}"
            ) from exc

        normalize_ooxml(target)
        self._validate(target)
        return GeneratedDocument(
            filename=filename,
            path=target,
            size_bytes=target.stat().st_size,
            type="word",
        )

    @staticmethod
    def _add_approval_header(doc, approval) -> None:
        """Render the approval-note metadata table and background."""
        header_rows = [
            ("Reference No.", approval.reference_number),
            ("Date", approval.date),
            ("Originator", approval.originator),
            ("Department", approval.department),
            ("Subject", approval.subject),
        ]
        header_table = doc.add_table(rows=len(header_rows), cols=2)
        header_table.style = "Table Grid"
        for row_index, (label, value) in enumerate(header_rows):
            header_table.cell(row_index, 0).text = label
            header_table.cell(row_index, 1).text = value

        if approval.background:
            doc.add_heading("Background", level=1)
            doc.add_paragraph(approval.background)

    @staticmethod
    def _add_approval_footer(doc, approval) -> None:
        """Render the recommendation and signature block that close the note."""
        if approval.recommendation:
            doc.add_heading("Recommendation", level=1)
            doc.add_paragraph(approval.recommendation)

        doc.add_heading("Approval", level=1)
        signatures = list(approval.signatures) or [ApprovalSignature()]
        signature_table = doc.add_table(rows=len(signatures) + 1, cols=4)
        signature_table.style = "Table Grid"
        for col_index, label in enumerate(("Name", "Designation", "Signature", "Date")):
            signature_table.cell(0, col_index).text = label
        for row_index, signature in enumerate(signatures, start=1):
            signature_table.cell(row_index, 0).text = signature.name
            signature_table.cell(row_index, 1).text = signature.designation
            signature_table.cell(row_index, 2).text = ""  # signed by hand
            signature_table.cell(row_index, 3).text = signature.date

    @staticmethod
    def _set_footer(doc, classification: str) -> None:
        """Write the footer: optional classification marking and "Page X of Y".

        A marketing line has no place on a formal approval note, so the footer
        carries the classification (when supplied) and live PAGE/NUMPAGES fields.
        """
        footer_paragraph = doc.sections[0].footer.paragraphs[0]
        footer_paragraph.text = ""
        if classification:
            classification_run = footer_paragraph.add_run(classification)
            classification_run.bold = True
            footer_paragraph.add_run("    |    ")
        footer_paragraph.add_run("Page ")
        WordDocumentGenerator._add_field(footer_paragraph, "PAGE")
        footer_paragraph.add_run(" of ")
        WordDocumentGenerator._add_field(footer_paragraph, "NUMPAGES")

    @staticmethod
    def _add_field(paragraph, field_name: str) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = field_name
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(instruction)
        run._r.append(end)

    @staticmethod
    def _add_section(doc, section: DocumentSection) -> None:
        if section.heading:
            doc.add_heading(section.heading, level=1)
        for paragraph in section.paragraphs:
            doc.add_paragraph(paragraph)
        for bullet in section.bullets:
            doc.add_paragraph(bullet, style="List Bullet")
        for item in section.numbered:
            doc.add_paragraph(item, style="List Number")
        table_rows = [row for row in section.table if row]
        if table_rows:
            rows = len(table_rows)
            cols = max(len(row) for row in table_rows)
            table = doc.add_table(rows=rows, cols=cols)
            table.style = "Table Grid"
            for row_index, row in enumerate(table_rows):
                for col_index in range(cols):
                    cell_text = row[col_index] if col_index < len(row) else ""
                    table.cell(row_index, col_index).text = cell_text
        for image in section.images:
            source = image.data or image.path
            if not source:
                continue
            from docx.shared import Inches

            width = Inches(image.width_inches or 6.0)
            # A non-web-safe source (bmp/gif/tiff/webp) or a PDF page rendered
            # for this figure arrives as PNG bytes; python-docx sniffs the
            # header from the stream.
            doc.add_picture(embed_source(source), width=width)
            if image.caption:
                caption_paragraph = doc.add_paragraph()
                caption_run = caption_paragraph.add_run(image.caption)
                caption_run.italic = True

    @staticmethod
    def _validate(path: Path) -> None:
        if not path.is_file() or path.stat().st_size <= 0:
            raise DocumentGenerationError("Generated document is missing or empty")
        try:
            from docx import Document

            Document(str(path))
        except Exception as exc:
            raise DocumentGenerationError(
                f"Generated document failed validation: {exc.__class__.__name__}"
            ) from exc


_XLSX_NUMBER_RE = re.compile(r"^-?(?:0|[1-9]\d*)(?:\.\d+)?$")
_INVALID_SHEET_CHARS_RE = re.compile(r"[\\/*?:\[\]]")


def _sheet_name(name: str, used: set[str]) -> str:
    """Return an Excel-legal, unique, <=31-char worksheet name."""
    base = _INVALID_SHEET_CHARS_RE.sub(" ", (name or "").strip())[:31].strip() or "Sheet"
    candidate = base
    index = 2
    while candidate.lower() in used:
        suffix = f" ({index})"
        candidate = base[: 31 - len(suffix)] + suffix
        index += 1
    used.add(candidate.lower())
    return candidate


def _coerce_cell(value: str):
    """Write '=...' as a real formula; keep other text verbatim.

    Plain integers/decimals are written as numbers so formulas can compute them,
    but values with leading zeros (IDs, codes) stay text.
    """
    if value.startswith("="):
        return value
    if _XLSX_NUMBER_RE.match(value):
        return float(value) if "." in value else int(value)
    return value


class XlsxDocumentGenerator(DocumentGenerator):
    """Generates valid .xlsx workbooks from structured content using openpyxl.

    Each section table becomes a worksheet (named from the section heading);
    cells beginning with ``=`` are written as real formulas. When there are no
    tables, section text is written as rows. Fully local; no macro/telemetry.
    """

    supported_types = ("excel",)

    async def generate(
        self, content: DocumentContent, output_dir: Path, filename: str
    ) -> GeneratedDocument:
        return await asyncio.to_thread(self._generate_sync, content, output_dir, filename)

    def _generate_sync(
        self, content: DocumentContent, output_dir: Path, filename: str
    ) -> GeneratedDocument:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise DocumentGenerationError(
                "openpyxl is not installed (Excel generation unavailable)"
            ) from exc

        output_dir.mkdir(parents=True, exist_ok=True)
        target = output_dir / filename

        try:
            workbook = Workbook()
            # Fixed timestamps keep byte-identical inputs deterministic.
            workbook.properties.created = datetime(1980, 1, 1)
            workbook.properties.modified = datetime(1980, 1, 1)
            used_names: set[str] = set()

            table_sections = [section for section in content.sections if any(section.table)]
            if table_sections:
                for index, section in enumerate(table_sections):
                    sheet = workbook.active if index == 0 else workbook.create_sheet()
                    sheet.title = _sheet_name(
                        section.heading or content.title or "Sheet", used_names
                    )
                    # A row with no cells widens nothing and only produces a
                    # degenerate grid; drop it rather than write a stray blank.
                    rows = [row for row in section.table if row]
                    for row_index, row in enumerate(rows, start=1):
                        for col_index, cell in enumerate(row, start=1):
                            target_cell = sheet.cell(row=row_index, column=col_index)
                            target_cell.value = _coerce_cell(str(cell))
                            value = target_cell.value
                            # Show measured values and numeric formulas (margin,
                            # MIN) to one decimal. Text formulas (IF/COUNTIF) keep
                            # the general format so they never show "2.0".
                            is_numeric_value = isinstance(value, float)
                            is_numeric_formula = (
                                isinstance(value, str)
                                and value.startswith("=")
                                and '"' not in value
                            )
                            if is_numeric_value or is_numeric_formula:
                                target_cell.number_format = "0.0"
                    column_count = max(len(row) for row in rows)
                    for col_index in range(1, column_count + 1):
                        sheet.cell(row=1, column=col_index).font = Font(bold=True)
                    # Freeze the header row and size columns to the widest text
                    # (formula strings are ignored so computed columns stay tidy).
                    sheet.freeze_panes = "A2"
                    for col_index in range(1, column_count + 1):
                        longest = 0
                        for row_index in range(1, len(rows) + 1):
                            value = sheet.cell(row=row_index, column=col_index).value
                            text = "" if value is None else str(value)
                            if text.startswith("="):
                                continue
                            longest = max(longest, len(text))
                        sheet.column_dimensions[get_column_letter(col_index)].width = min(
                            max(longest + 2, 8), 30
                        )
                    # Images sit one blank column clear of the grid, so no
                    # figure overlaps the data it illustrates.
                    self._add_images(sheet, section, column_count + 2, 2)
            else:
                sheet = workbook.active
                sheet.title = _sheet_name(content.title or "Sheet", used_names)
                row_index = 1
                for section in content.sections:
                    if section.heading:
                        heading_cell = sheet.cell(row=row_index, column=1)
                        heading_cell.value = section.heading
                        heading_cell.font = Font(bold=True)
                        row_index += 1
                    for item in (*section.paragraphs, *section.bullets, *section.numbered):
                        sheet.cell(row=row_index, column=1).value = item
                        row_index += 1
                    # Column B keeps the picture clear of the text in column A.
                    self._add_images(sheet, section, 2, row_index)

            if content.sources:
                sheet = workbook.create_sheet(title=_sheet_name("Sources", used_names))
                reference_header = sheet.cell(row=1, column=1)
                reference_header.value = "Reference"
                reference_header.font = Font(bold=True)
                for row_index, source in enumerate(content.sources, start=2):
                    sheet.cell(row=row_index, column=1).value = source
                sheet.column_dimensions["A"].width = 30
                sheet.freeze_panes = "A2"

            workbook.save(str(target))
        except DocumentGenerationError:
            raise
        except ImageNormalizationError as exc:
            raise DocumentGenerationError(f"Excel generation failed: {exc}") from exc
        except Exception as exc:
            raise DocumentGenerationError(
                f"Excel generation failed: {exc.__class__.__name__}"
            ) from exc

        normalize_ooxml(target)
        self._validate(target)
        return GeneratedDocument(
            filename=filename,
            path=target,
            size_bytes=target.stat().st_size,
            type="excel",
        )

    @staticmethod
    def _add_images(sheet, section: DocumentSection, anchor_col: int, start_row: int) -> None:
        """Embed a section's images below each other in one blank column.

        Each picture is scaled to ``width_inches`` (default 6in) with its aspect
        ratio preserved. A caption is written into the cell under the picture,
        since a worksheet has no caption primitive. Rows are counted in the
        20-pixel rows used for anchoring, so two pictures never overlap.
        """
        if not section.images:
            return
        from openpyxl.drawing.image import Image as XlsxImage
        from openpyxl.utils import get_column_letter

        row = max(start_row, 1)
        for image in section.images:
            source = image.data or image.path
            if not source:
                continue
            picture = XlsxImage(embed_source(source))
            natural_width, natural_height = picture.width, picture.height
            target_width = int((image.width_inches or 6.0) * 96)
            if natural_width and target_width > 0:
                picture.height = int(natural_height * (target_width / natural_width))
                picture.width = target_width
            picture.anchor = f"{get_column_letter(anchor_col)}{row}"
            sheet.add_image(picture)
            row += max(2, int(picture.height / 20) + 1)
            if image.caption:
                sheet.cell(row=row, column=anchor_col).value = image.caption
                row += 1

    @staticmethod
    def _validate(path: Path) -> None:
        if not path.is_file() or path.stat().st_size <= 0:
            raise DocumentGenerationError("Generated workbook is missing or empty")
        try:
            from openpyxl import load_workbook

            load_workbook(str(path))
        except Exception as exc:
            raise DocumentGenerationError(
                f"Generated workbook failed validation: {exc.__class__.__name__}"
            ) from exc
