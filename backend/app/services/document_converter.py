"""Local document format conversion service (PDF, DOCX, PPTX).

Fully local, zero-telemetry conversion pipeline supporting:
- PDF -> Word (DOCX)
- Word (DOCX) -> PDF
- PowerPoint (PPTX) -> PDF

All transformations run locally in memory/disk using python-docx, pypdf/pypdfium2,
python-pptx, and ReportLab. No external cloud converters or egress network calls.
"""

from __future__ import annotations

import html
import io
import logging
from pathlib import Path
from typing import Optional

logger = logging.getLogger("app.services.document_converter")


class DocumentConversionError(Exception):
    """Raised when document conversion fails."""


def convert_pdf_to_docx(pdf_path: Path, output_path: Path) -> Path:
    """Convert a PDF file to a Microsoft Word (.docx) document."""
    try:
        import pypdf
        import docx
        from docx.shared import Inches, Pt, RGBColor

        reader = pypdf.PdfReader(str(pdf_path))
        doc = docx.Document()

        # Set clean document properties
        doc.core_properties.title = pdf_path.stem.replace("_", " ").title()
        doc.core_properties.author = "AstraSovereign Local Conversion Engine"

        # Add title header
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(pdf_path.stem.replace("_", " ").title())
        title_run.bold = True
        title_run.font.size = Pt(18)
        title_run.font.color.rgb = RGBColor(17, 24, 39)

        meta_p = doc.add_paragraph()
        meta_run = meta_p.add_run(f"Converted from {pdf_path.name} • AstraSovereign Air-Gapped Engine")
        meta_run.italic = True
        meta_run.font.size = Pt(9.5)
        meta_run.font.color.rgb = RGBColor(107, 114, 128)

        total_pages = len(reader.pages)
        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            lines = [line.strip() for line in text.splitlines() if line.strip()]

            if total_pages > 1 and page_idx > 0:
                doc.add_page_break()
                p_header = doc.add_paragraph()
                p_run = p_header.add_run(f"— Page {page_idx + 1} —")
                p_run.font.size = Pt(9)
                p_run.font.color.rgb = RGBColor(156, 163, 175)

            current_p = None
            for line in lines:
                # Detect heading-like short lines
                if len(line) < 60 and (line.isupper() or line.endswith(":") or (line[0].isdigit() and "." in line[:4])):
                    h = doc.add_paragraph()
                    hr = h.add_run(line)
                    hr.bold = True
                    hr.font.size = Pt(12)
                    hr.font.color.rgb = RGBColor(31, 41, 55)
                    current_p = None
                else:
                    if current_p is None:
                        current_p = doc.add_paragraph()
                    else:
                        current_p.add_run(" ")
                    run = current_p.add_run(line)
                    run.font.size = Pt(10.5)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path
    except Exception as exc:
        logger.exception("PDF to DOCX conversion failed: %s", exc)
        raise DocumentConversionError(f"Failed to convert PDF to DOCX: {exc}") from exc


def convert_docx_to_pdf(docx_path: Path, output_path: Path) -> Path:
    """Convert a Microsoft Word (.docx) document to a PDF file using ReportLab."""
    try:
        import docx
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib import colors
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak

        doc = docx.Document(str(docx_path))
        pdf = SimpleDocTemplate(
            str(output_path),
            pagesize=letter,
            rightMargin=54,
            leftMargin=54,
            topMargin=54,
            bottomMargin=54,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Title"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#0f172a"),
            spaceAfter=12,
            alignment=0,
        )
        heading_style = ParagraphStyle(
            "DocHeading",
            parent=styles["Heading2"],
            fontSize=13,
            leading=16,
            textColor=colors.HexColor("#1e293b"),
            spaceBefore=10,
            spaceAfter=6,
            keepWithNext=True,
        )
        body_style = ParagraphStyle(
            "DocBody",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#334155"),
            spaceAfter=6,
        )
        meta_style = ParagraphStyle(
            "DocMeta",
            parent=styles["Normal"],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#64748b"),
            spaceAfter=14,
        )

        story = []

        # Add Title & Header
        title_text = html.escape(docx_path.stem.replace("_", " ").title())
        story.append(Paragraph(title_text, title_style))
        story.append(
            Paragraph(
                f"Converted from {html.escape(docx_path.name)} • AstraSovereign Air-Gapped Engine",
                meta_style,
            )
        )
        story.append(Spacer(1, 10))

        for p in doc.paragraphs:
            raw_text = p.text.strip()
            if not raw_text:
                continue

            escaped = html.escape(raw_text)
            # Detect heading style or attributes
            if p.style.name.startswith("Heading") or (len(raw_text) < 60 and raw_text.isupper()):
                story.append(Paragraph(escaped, heading_style))
            else:
                story.append(Paragraph(escaped, body_style))

        # Convert tables if any
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = [html.escape(cell.text.strip()) for cell in row.cells]
                table_data.append([Paragraph(c or "-", body_style) for c in row_data])

            if table_data:
                t = Table(table_data)
                t.setStyle(
                    TableStyle([
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f8fafc")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                    ])
                )
                story.append(Spacer(1, 8))
                story.append(t)
                story.append(Spacer(1, 10))

        output_path.parent.mkdir(parents=True, exist_ok=True)
        pdf.build(story)
        return output_path
    except Exception as exc:
        logger.exception("DOCX to PDF conversion failed: %s", exc)
        raise DocumentConversionError(f"Failed to convert DOCX to PDF: {exc}") from exc


def convert_pptx_to_pdf(pptx_path: Path, output_path: Path) -> Path:
    """Convert a PowerPoint (.pptx) presentation to a landscape PDF slide deck."""
    try:
        import pptx
        from reportlab.lib.pagesizes import landscape, letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib import colors
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, PageBreak, Table, TableStyle

        prs = pptx.Presentation(str(pptx_path))
        pdf = SimpleDocTemplate(
            str(output_path),
            pagesize=landscape(letter),
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40,
        )

        styles = getSampleStyleSheet()
        slide_title_style = ParagraphStyle(
            "SlideTitle",
            parent=styles["Title"],
            fontSize=22,
            leading=26,
            textColor=colors.HexColor("#0f172a"),
            spaceAfter=14,
            alignment=0,
        )
        slide_num_style = ParagraphStyle(
            "SlideNum",
            parent=styles["Normal"],
            fontSize=9,
            leading=11,
            textColor=colors.HexColor("#ef4444"),
            spaceAfter=4,
            alignment=0,
        )
        bullet_style = ParagraphStyle(
            "SlideBullet",
            parent=styles["Normal"],
            fontSize=12,
            leading=17,
            textColor=colors.HexColor("#1e293b"),
            leftIndent=15,
            spaceAfter=8,
        )
        body_style = ParagraphStyle(
            "SlideBody",
            parent=styles["Normal"],
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#334155"),
            spaceAfter=6,
        )

        story = []
        total_slides = len(prs.slides)

        for s_idx, slide in enumerate(prs.slides):
            if s_idx > 0:
                story.append(PageBreak())

            story.append(Paragraph(f"SLIDE {s_idx + 1} OF {total_slides}", slide_num_style))

            # Extract title if present
            slide_title = ""
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                slide_title = slide.shapes.title.text_frame.text.strip()

            if slide_title:
                story.append(Paragraph(html.escape(slide_title), slide_title_style))
            else:
                story.append(Paragraph(f"Slide {s_idx + 1}", slide_title_style))

            story.append(Spacer(1, 10))

            # Extract shapes & textframes
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        text = p.text.strip()
                        if not text:
                            continue
                        escaped = html.escape(text)
                        if p.level and p.level > 0 or text.startswith("-") or text.startswith("•"):
                            clean_t = text.lstrip("-•").strip()
                            story.append(Paragraph(f"&bull; {html.escape(clean_t)}", bullet_style))
                        else:
                            story.append(Paragraph(escaped, body_style))

                elif shape.has_table:
                    table_data = []
                    for row in shape.table.rows:
                        row_data = [html.escape(cell.text.strip()) for cell in row.cells]
                        table_data.append([Paragraph(c or "-", body_style) for c in row_data])
                    if table_data:
                        t = Table(table_data)
                        t.setStyle(
                            TableStyle([
                                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                                ("TOPPADDING", (0, 0), (-1, -1), 6),
                                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                            ])
                        )
                        story.append(Spacer(1, 6))
                        story.append(t)
                        story.append(Spacer(1, 8))

        output_path.parent.mkdir(parents=True, exist_ok=True)
        pdf.build(story)
        return output_path
    except Exception as exc:
        logger.exception("PPTX to PDF conversion failed: %s", exc)
        raise DocumentConversionError(f"Failed to convert PPTX to PDF: {exc}") from exc
