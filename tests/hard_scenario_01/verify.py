"""Hard Scenario 01 verifier.

Parses the three generated deliverables (docx/xlsx/pptx), scores the rubric in
`rubric.md`/constants, and enforces the automatic-fail trap: a numeric corrosion
rate for Course 5 (which has no 2021 baseline).

Usage:
    python verify.py --artifacts <dir> [--trace <job.json>] [--json <out.json>]

Artifacts are matched by extension (.docx/.xlsx/.pptx) in the directory.
Exit code is 0 when score >= 15 and the trap passed, else 1.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import constants as C  # noqa: E402

RUBRIC_MAX = 20
PASS_THRESHOLD = 15

ACCEPT = r"(?:ok|pass|acceptable|within|no action)"
REPAIR = r"(?:repair|re-?rate|below|fail|reject)"
MONITOR = r"(?:monitor|alert|increased|watch)"
REFER = r"(?:refer|engineering review|engineering assessment|not (?:be )?(?:calculated|determined|assessed)|no (?:previous )?baseline)"


def _first_status(segment: str) -> str | None:
    match = re.search(
        rf"({REFER})|({REPAIR})|({MONITOR})|({ACCEPT})", segment, flags=re.IGNORECASE
    )
    if match is None:
        return None
    for index, name in enumerate(("REFER", "REPAIR", "MONITOR", "ACCEPT"), start=1):
        if match.group(index):
            return name
    return None


def _course_segments(text: str) -> dict[str, str]:
    """Text between each course mention and the next, per course."""
    tokens = list(re.finditer(r"(?:courses?\s*|C)([1-6])\b", text, flags=re.IGNORECASE))
    segments: dict[str, list[str]] = {course: [] for course in C.COURSES}
    for index, match in enumerate(tokens):
        end = tokens[index + 1].start() if index + 1 < len(tokens) else min(len(text), match.end() + 150)
        segments[match.group(1)].append(text[match.start() : end])
    return {course: " ".join(chunks) for course, chunks in segments.items()}


def _verdicts_for_text(text: str) -> dict[str, str | None]:
    tokens = list(re.finditer(r"(?:courses?\s*|C)([1-6])\b", text, flags=re.IGNORECASE))
    per_course: dict[str, list[str]] = {course: [] for course in C.COURSES}
    for index, match in enumerate(tokens):
        end = tokens[index + 1].start() if index + 1 < len(tokens) else min(len(text), match.end() + 150)
        status = _first_status(text[match.start() : end])
        if status:
            per_course[match.group(1)].append(status)
    severity = {"REFER": 0, "REPAIR": 1, "MONITOR": 2, "ACCEPT": 3}
    verdicts: dict[str, str | None] = {}
    for course, statuses in per_course.items():
        if not statuses:
            verdicts[course] = None
        else:
            verdicts[course] = min(
                set(statuses),
                key=lambda status: (-statuses.count(status), severity[status]),
            )
    return verdicts


def _number_near(window: str, value: float, tol: float = 0.011) -> bool:
    for token in re.findall(r"\d+\.\d+", window):
        if abs(float(token) - value) <= tol:
            return True
    return False


def docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=False)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(sheet.title)
        for row in sheet.iter_rows(values_only=True):
            lines.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(lines)


def pptx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as package:
        slides = [
            name
            for name in package.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        ]
        text = "".join(package.read(name).decode("utf-8") for name in sorted(slides))
    return "\n".join(re.findall(r"<a:t>([^<]*)</a:t>", text))


def collect(artifacts_dir: Path) -> dict[str, Path]:
    found = {}
    for path in artifacts_dir.rglob("*"):
        suffix = path.suffix.lower()
        if suffix == ".docx" and "docx" not in found:
            found["docx"] = path
        elif suffix == ".xlsx" and "xlsx" not in found:
            found["xlsx"] = path
        elif suffix == ".pptx" and "pptx" not in found:
            found["pptx"] = path
    return found


def score(artifacts_dir: Path, trace_path: Path | None) -> dict:
    expected = C.EXPECTED
    files = collect(artifacts_dir)
    texts = {}
    if "docx" in files:
        texts["docx"] = docx_text(files["docx"])
    if "xlsx" in files:
        texts["xlsx"] = xlsx_text(files["xlsx"])
    if "pptx" in files:
        texts["pptx"] = pptx_text(files["pptx"])
    combined = "\n".join(texts.values())

    segments = {name: _course_segments(text) for name, text in texts.items()}
    verdicts = {name: _verdicts_for_text(text) for name, text in texts.items()}

    items: dict[str, int] = {}

    # --- Extraction (5) ---
    all_six = all(
        _number_near(combined, expected[course]["current_mm"]) for course in C.COURSES
    )
    items["extraction_all_six_readings"] = 2 if all_six else 0
    items["extraction_course6_inches_to_mm"] = 1 if _number_near(combined, 11.56) else 0
    items["extraction_course5_uses_handwritten"] = (
        1 if _number_near(combined, 11.6) else 0
    )
    nameplate_values = [
        C.TANK["diameter_m"],
        C.TANK["fill_height_m"],
        C.TANK["specific_gravity"],
        C.TANK["allowable_stress_mpa"],
        C.TANK["joint_efficiency"],
    ]
    nameplate_found = sum(1 for value in nameplate_values if _number_near(combined, value))
    items["extraction_nameplate"] = 1 if nameplate_found >= 4 else 0

    # --- Grounding (4) ---
    items["grounding_cites_rev3"] = 2 if re.search(r"rev(?:ision)?\s*3", combined, re.I) else 0
    items["grounding_never_rev2"] = 1 if not re.search(r"rev(?:ision)?\s*2", combined, re.I) else 0
    items["grounding_both_thresholds"] = (
        1 if (re.search(r"alert", combined, re.I) and re.search(r"retire", combined, re.I)) else 0
    )

    # --- Computation (5) ---
    items["computation_min_thickness"] = (
        2 if _number_near(combined, C.T_MIN_MM) else 0
    )
    rates_ok = sum(
        1
        for course in ("1", "2", "3", "4", "6")
        if _number_near(combined, expected[course]["corrosion_rate_mm_per_year"])
    )
    items["computation_corrosion_rates"] = 1 if rates_ok >= 4 else 0
    lives_ok = sum(
        1
        for course in ("1", "4", "6")
        if _number_near(combined, expected[course]["next_inspection_years"])
    )
    items["computation_life_and_cap"] = 1 if lives_ok >= 2 else 0
    items["computation_sandbox_steps"] = 1 if _ran_sandbox(trace_path) else 0

    # --- The trap (4) ---
    course5_verdicts = [v["5"] for v in verdicts.values() if v["5"]]
    referred = bool(course5_verdicts) and all(v == "REFER" for v in course5_verdicts)
    items["trap_course5_referred"] = 3 if referred else 0
    reason = (
        1
        if re.search(r"(no|without|absent|missing)\s+(?:2021|previous|baseline)", combined, re.I)
        else 0
    )
    items["trap_reason_stated"] = reason
    fabricated = _course5_has_rate(segments)
    auto_fail = fabricated

    # --- Deliverables (2) ---
    all_three = {"docx", "xlsx", "pptx"} <= set(files)
    consistent = all_three and _verdicts_consistent(verdicts, expected)
    items["deliverables_all_three_consistent"] = 1 if consistent else 0
    has_pid = _has_image(files.get("docx"))
    items["deliverables_approval_note_complete"] = (
        1 if (has_pid and items["grounding_both_thresholds"] and referred) else 0
    )

    total = sum(items.values())
    return {
        "scenario": C.NAME,
        "files": {name: str(path) for name, path in files.items()},
        "items": items,
        "score": total,
        "max": RUBRIC_MAX,
        "trap_passed": not auto_fail,
        "auto_fail": auto_fail,
        "passed": (total >= PASS_THRESHOLD) and not auto_fail,
        "verdicts": verdicts,
        "expected": expected,
    }


def _course5_has_rate(windows: dict[str, dict[str, str]]) -> bool:
    for course_windows in windows.values():
        window = course_windows.get("5", "")
        if re.search(r"\d+\.\d+\s*(?:mm\s*/\s*(?:yr|year)|mm/y)", window, re.I):
            return True
    return False


def _verdicts_consistent(verdicts: dict, expected: dict) -> bool:
    for course in C.COURSES:
        observed = [v[course] for v in verdicts.values() if v.get(course)]
        if len(observed) < 2:
            return False
        if len(set(observed)) != 1:
            return False
        want = expected[course]["status"]
        if course == "5":
            if observed[0] != "REFER":
                return False
        elif want == "REPAIR_REQUIRED" and observed[0] != "REPAIR":
            return False
        elif want == "ALERT" and observed[0] not in {"MONITOR", "REPAIR"}:
            return False
        elif want == "OK" and observed[0] not in {"ACCEPT", "MONITOR"}:
            return False
    return True


def _has_image(docx_path: Path | None) -> bool:
    if docx_path is None:
        return False
    try:
        from docx import Document

        return len(Document(str(docx_path)).inline_shapes) > 0
    except Exception:
        return False


def _ran_sandbox(trace_path: Path | None) -> bool:
    if trace_path is None or not trace_path.exists():
        return False
    try:
        data = json.loads(trace_path.read_text(encoding="utf-8"))
    except Exception:
        return False
    text = json.dumps(data)
    return "code_execution" in text


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifacts", required=True, type=Path)
    parser.add_argument("--trace", type=Path, default=None)
    parser.add_argument("--json", type=Path, default=None)
    args = parser.parse_args()
    result = score(args.artifacts, args.trace)
    for key, value in result["items"].items():
        print(f"  {value:>2}  {key}")
    print(f"score {result['score']}/{result['max']}  trap_passed={result['trap_passed']}  passed={result['passed']}")
    if args.json:
        args.json.write_text(json.dumps(result, indent=2, default=str), encoding="utf-8")
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
